# church_offering_gui_improved.py
# -*- coding: utf-8 -*-

import sqlite3
from datetime import datetime, date, timedelta
from typing import Optional, List, Tuple, Dict
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import shutil

# PDF 생성을 위한 라이브러리
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

DB_PATH = "church_offerings.db"

# =========================
# DB Functions
# =========================
def connect_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def init_db() -> None:
    conn = connect_db()
    cur = conn.cursor()

    cur.execute("""
    CREATE TABLE IF NOT EXISTS members(
        code TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS offering_types(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL UNIQUE,
        display_order INTEGER DEFAULT 0
    );
    """)

    # 기존 테이블에 컬럼 추가 (마이그레이션)
    try:
        cur.execute("SELECT display_order FROM offering_types LIMIT 1;")
    except sqlite3.OperationalError:
        # display_order 컬럼이 없으면 추가
        cur.execute("ALTER TABLE offering_types ADD COLUMN display_order INTEGER DEFAULT 0;")

    # memo 컬럼 추가 (마이그레이션)
    try:
        cur.execute("SELECT memo FROM offerings LIMIT 1;")
    except sqlite3.OperationalError:
        # memo 컬럼이 없으면 추가
        cur.execute("ALTER TABLE offerings ADD COLUMN memo TEXT;")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS offerings(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        offering_date TEXT NOT NULL,
        offering_type TEXT NOT NULL,
        member_code TEXT NOT NULL,
        amount INTEGER NOT NULL CHECK(amount >= 0),
        memo TEXT,
        created_at TEXT NOT NULL,
        updated_at TEXT,
        FOREIGN KEY(member_code) REFERENCES members(code) ON UPDATE CASCADE ON DELETE RESTRICT
    );
    """)

    cur.execute("CREATE INDEX IF NOT EXISTS idx_offerings_date ON offerings(offering_date);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_offerings_type ON offerings(offering_type);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_offerings_member ON offerings(member_code);")

    cur.execute("SELECT DISTINCT offering_type FROM offerings;")
    existing = [r[0] for r in cur.fetchall()]
    for t in existing:
        if t and isinstance(t, str):
            cur.execute("INSERT OR IGNORE INTO offering_types(name) VALUES(?)", (t.strip(),))

    conn.commit()
    conn.close()

def backup_database() -> str:
    """데이터베이스 백업"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = f"church_offerings_backup_{timestamp}.db"
    shutil.copy2(DB_PATH, backup_path)
    return backup_path

def normalize_date(date_str: str) -> str:
    dt = datetime.strptime(date_str.strip(), "%Y-%m-%d")
    return dt.strftime("%Y-%m-%d")

# ---- offering types ----
def get_offering_types() -> List[str]:
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("SELECT name FROM offering_types ORDER BY display_order, name;")
    rows = cur.fetchall()
    conn.close()
    return [r[0] for r in rows]

def add_offering_type(name: str) -> None:
    name = name.strip()
    if not name:
        raise ValueError("빈 이름은 추가할 수 없어.")
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("SELECT COALESCE(MAX(display_order), 0) + 1 FROM offering_types")
    order = cur.fetchone()[0]
    cur.execute("INSERT INTO offering_types(name, display_order) VALUES(?, ?)", (name, order))
    conn.commit()
    conn.close()

def rename_offering_type(old_name: str, new_name: str) -> None:
    old_name = old_name.strip()
    new_name = new_name.strip()
    if not old_name or not new_name:
        raise ValueError("이름이 비었어.")

    conn = connect_db()
    cur = conn.cursor()

    cur.execute("SELECT 1 FROM offering_types WHERE name=?", (new_name,))
    if cur.fetchone():
        conn.close()
        raise ValueError("이미 존재하는 헌금 종류 이름이야.")

    cur.execute("UPDATE offering_types SET name=? WHERE name=?", (new_name, old_name))
    if cur.rowcount == 0:
        conn.close()
        raise ValueError("변경할 기존 헌금 종류를 찾지 못했어.")

    cur.execute("UPDATE offerings SET offering_type=? WHERE offering_type=?", (new_name, old_name))
    conn.commit()
    conn.close()

def delete_offering_type(name: str) -> None:
    name = name.strip()
    if not name:
        return

    conn = connect_db()
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM offerings WHERE offering_type=?", (name,))
    cnt = int(cur.fetchone()[0] or 0)
    if cnt > 0:
        conn.close()
        raise ValueError("이미 입력된 헌금 내역이 있어서 삭제할 수 없어.")

    cur.execute("DELETE FROM offering_types WHERE name=?", (name,))
    conn.commit()
    conn.close()

# ---- members ----
def get_member_name(code: str) -> Optional[str]:
    code = code.strip()
    if not code:
        return None
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("SELECT name FROM members WHERE code=?", (code,))
    row = cur.fetchone()
    conn.close()
    return row[0] if row else None

def upsert_member(code: str, name: str) -> None:
    code = code.strip()
    name = name.strip()
    if not code or not name:
        return
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    INSERT INTO members(code, name) VALUES(?, ?)
    ON CONFLICT(code) DO UPDATE SET name = excluded.name;
    """, (code, name))
    conn.commit()
    conn.close()

def delete_member(code: str) -> None:
    code = code.strip()
    if not code:
        return
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM members WHERE code=?", (code,))
    conn.commit()
    conn.close()

def query_members(keyword: str = "") -> List[Tuple[str, str]]:
    kw = (keyword or "").strip()
    conn = connect_db()
    cur = conn.cursor()

    if kw:
        like = f"%{kw}%"
        cur.execute("""
        SELECT code, name
        FROM members
        WHERE name LIKE ? OR code LIKE ?
        ORDER BY name, code;
        """, (like, like))
    else:
        cur.execute("SELECT code, name FROM members ORDER BY name, code;")

    rows = cur.fetchall()
    conn.close()
    return rows

def get_member_code_by_name_exact(name: str) -> Optional[str]:
    name = name.strip()
    if not name:
        return None
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("SELECT code FROM members WHERE name=? ORDER BY code LIMIT 1;", (name,))
    row = cur.fetchone()
    conn.close()
    return row[0] if row else None

def ensure_member_code_by_name(name: str) -> str:
    """코드가 없으면 9 + 8자리 숫자 임시코드 생성"""
    name = name.strip()
    code = get_member_code_by_name_exact(name)
    if code:
        return code

    conn = connect_db()
    cur = conn.cursor()

    cur.execute("SELECT code FROM members WHERE code LIKE '9________' ORDER BY code DESC LIMIT 1;")
    row = cur.fetchone()

    if row:
        try:
            last_num = int(row[0])
            temp_code = str(last_num + 1)
        except:
            temp_code = f"9{datetime.now().strftime('%Y%m%d')}"
    else:
        temp_code = f"9{datetime.now().strftime('%Y%m%d')}"

    conn.close()
    upsert_member(temp_code, name)
    return temp_code

def search_members_with_stats(keyword: str, limit: int = 200) -> List[Tuple[str, str, str, int]]:
    kw = (keyword or "").strip()
    if not kw:
        return []

    like = f"%{kw}%"
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT
        m.code,
        m.name,
        COALESCE(MAX(o.offering_date), '') AS last_date,
        COALESCE(SUM(o.amount), 0) AS total_amount
    FROM members m
    LEFT JOIN offerings o ON o.member_code = m.code
    WHERE m.name LIKE ?
    GROUP BY m.code, m.name
    ORDER BY m.name, m.code
    LIMIT ?;
    """, (like, int(limit)))
    rows = cur.fetchall()
    conn.close()
    return [(str(c), str(n), str(ld or ""), int(t or 0)) for (c, n, ld, t) in rows]

def import_members_from_excel(filepath: str) -> Tuple[int, int]:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.worksheets[0]

    header1 = str(ws.cell(row=1, column=1).value or "").strip().lower()
    header2 = str(ws.cell(row=1, column=2).value or "").strip().lower()

    start_row = 2 if (
        header1 in ("code", "member_code", "사용자코드", "코드") and
        header2 in ("name", "username", "이름")
    ) else 1

    ok = 0
    skip = 0

    for r in range(start_row, ws.max_row + 1):
        code = str(ws.cell(row=r, column=1).value or "").strip()
        name = str(ws.cell(row=r, column=2).value or "").strip()
        if not code or not name:
            skip += 1
            continue
        upsert_member(code, name)
        ok += 1

    return ok, skip

# ---- offerings ----
def insert_offering(date_: str, type_: str, code: str, amount: int, memo: str = "") -> None:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    INSERT INTO offerings(offering_date, offering_type, member_code, amount, memo, created_at)
    VALUES(?,?,?,?,?,?)
    """, (date_, type_, code.strip(), int(amount), memo.strip(), now))
    conn.commit()
    conn.close()

    conn = connect_db()
    cur = conn.cursor()
    cur.execute("INSERT OR IGNORE INTO offering_types(name) VALUES(?)", (type_.strip(),))
    conn.commit()
    conn.close()

def update_offering(oid: int, date_: str, type_: str, code: str, amount: int, memo: str = "") -> None:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    UPDATE offerings
    SET offering_date=?, offering_type=?, member_code=?, amount=?, memo=?, updated_at=?
    WHERE id=?
    """, (date_, type_, code.strip(), int(amount), memo.strip(), now, int(oid)))
    conn.commit()
    conn.close()

    conn = connect_db()
    cur = conn.cursor()
    cur.execute("INSERT OR IGNORE INTO offering_types(name) VALUES(?)", (type_.strip(),))
    conn.commit()
    conn.close()

def delete_offering(oid: int) -> None:
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM offerings WHERE id=?", (int(oid),))
    conn.commit()
    conn.close()

def delete_offerings_batch(oid_list: List[int]) -> int:
    """여러 헌금 내역 한번에 삭제"""
    if not oid_list:
        return 0
    conn = connect_db()
    cur = conn.cursor()
    placeholders = ','.join('?' * len(oid_list))
    cur.execute(f"DELETE FROM offerings WHERE id IN ({placeholders})", oid_list)
    deleted = cur.rowcount
    conn.commit()
    conn.close()
    return deleted

def query_offerings_by_range(start_date: str, end_date: str) -> List[Tuple]:
    s = normalize_date(start_date)
    e = normalize_date(end_date)
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT o.id, o.offering_date, o.offering_type, o.member_code, m.name, o.amount, COALESCE(o.memo, '')
    FROM offerings o
    JOIN members m ON m.code = o.member_code
    WHERE o.offering_date BETWEEN ? AND ?
    ORDER BY o.offering_date DESC, o.id DESC;
    """, (s, e))
    rows = cur.fetchall()
    conn.close()
    return rows

def sum_total_by_range(start_date: str, end_date: str) -> int:
    s = normalize_date(start_date)
    e = normalize_date(end_date)
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT COALESCE(SUM(amount),0)
    FROM offerings
    WHERE offering_date BETWEEN ? AND ?;
    """, (s, e))
    val = int(cur.fetchone()[0] or 0)
    conn.close()
    return val

def sum_by_type_by_range(start_date: str, end_date: str) -> Dict[str, int]:
    s = normalize_date(start_date)
    e = normalize_date(end_date)
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT offering_type, COALESCE(SUM(amount),0)
    FROM offerings
    WHERE offering_date BETWEEN ? AND ?
    GROUP BY offering_type
    ORDER BY offering_type;
    """, (s, e))
    rows = cur.fetchall()
    conn.close()
    return {str(t): int(v or 0) for (t, v) in rows}

def export_pdf_report(filepath: str, start_date: str, end_date: str,
                      check_data: Dict, expense_amt: int, expense_memo: str) -> None:
    """PDF 리포트 생성 (1페이지)"""
    if not PDF_AVAILABLE:
        raise ImportError("reportlab 라이브러리가 설치되어 있지 않습니다.")

    s = normalize_date(start_date)
    e = normalize_date(end_date)

    conn = connect_db()
    cur = conn.cursor()

    # 전체 합계
    cur.execute("""
    SELECT COALESCE(SUM(amount), 0)
    FROM offerings
    WHERE offering_date BETWEEN ? AND ?;
    """, (s, e))
    grand_total = int(cur.fetchone()[0] or 0)

    # 헌금 종류별 합계
    cur.execute("""
    SELECT offering_type, COALESCE(SUM(amount), 0)
    FROM offerings
    WHERE offering_date BETWEEN ? AND ?
    GROUP BY offering_type
    ORDER BY offering_type;
    """, (s, e))
    type_totals = cur.fetchall()

    conn.close()

    # PDF 문서 생성 (여백 최소화)
    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        topMargin=10*mm,
        bottomMargin=10*mm,
        leftMargin=15*mm,
        rightMargin=15*mm
    )
    elements = []

    # 한글 폰트 등록
    font_name = 'Helvetica'
    font_paths = [
        'C:/Windows/Fonts/malgun.ttf',
        'C:/Windows/Fonts/gulim.ttc',
        'C:/Windows/Fonts/batang.ttc',
        'C:/Windows/Fonts/NanumGothic.ttf',
    ]

    for font_path in font_paths:
        try:
            import os
            if os.path.exists(font_path):
                pdfmetrics.registerFont(TTFont('KoreanFont', font_path))
                font_name = 'KoreanFont'
                break
        except:
            continue

    # 스타일 정의 (간격 축소)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontName=font_name,
        fontSize=16,
        textColor=colors.HexColor('#2C3E50'),
        spaceAfter=8,
        alignment=TA_CENTER
    )

    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontName=font_name,
        fontSize=11,
        textColor=colors.HexColor('#34495E'),
        spaceAfter=4,
        spaceBefore=6
    )

    # 제목
    title = Paragraph(f"헌금 대조 보고서 ({s} ~ {e})", title_style)
    elements.append(title)
    elements.append(Spacer(1, 3*mm))

    # 상단 요약 (기간 헌금 + 종류별 합계를 가로로 배치)
    left_data = [
        ['기간 헌금 합계'],
        [f'{grand_total:,}원']
    ]

    left_table = Table(left_data, colWidths=[55*mm])
    left_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498DB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('FONTSIZE', (0, 1), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))

    # 헌금 종류별 (최대 5개만)
    right_data = [['헌금 종류', '금액']]
    for i, (t_name, t_sum) in enumerate(type_totals[:5]):
        right_data.append([t_name, f'{int(t_sum):,}원'])

    right_table = Table(right_data, colWidths=[60*mm, 45*mm])
    right_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#27AE60')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))

    # 상단을 가로로 배치
    top_row = Table([[left_table, right_table]], colWidths=[60*mm, 110*mm])
    top_row.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(top_row)
    elements.append(Spacer(1, 4*mm))

    # 실물 대조 (컴팩트)
    section_title = Paragraph("실물 대조", heading_style)
    elements.append(section_title)

    check_detail_data = [['항목', '단가', '수량', '금액']]
    check_total = 0

    for item_name, item_data in check_data.items():
        qty = item_data.get('qty', 0)
        amt = item_data.get('amt', 0)
        unit = item_data.get('unit', 0)

        if unit > 0:
            item_total = qty * unit if qty > 0 else amt
        else:
            item_total = amt

        check_total += item_total

        if item_total > 0:
            check_detail_data.append([
                item_name,
                f'{unit:,}' if unit > 0 else '-',
                str(qty) if qty > 0 else '-',
                f'{item_total:,}원'
            ])

    check_table = Table(check_detail_data, colWidths=[35*mm, 30*mm, 20*mm, 35*mm])
    check_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E67E22')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))
    elements.append(check_table)
    elements.append(Spacer(1, 2*mm))

    # 하단 정산 (가로 배치)
    bottom_data = [
        ['대조 합계', f'{check_total:,}원'],
        ['지출 금액', f'{expense_amt:,}원'],
        ['실제 잔액', f'{check_total - expense_amt:,}원'],
        ['차이', f'{grand_total - (check_total - expense_amt):,}원']
    ]

    # 지출 메모가 있으면 추가
    if expense_memo:
        memo_short = expense_memo[:30] + '...' if len(expense_memo) > 30 else expense_memo
        bottom_data.insert(2, ['사용처', memo_short])

    bottom_table = Table(bottom_data, colWidths=[50*mm, 50*mm])
    bottom_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#34495E')),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E74C3C')),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.whitesmoke),
    ]))

    # 오른쪽 정렬
    bottom_wrapper = Table([[bottom_table]], colWidths=[100*mm], hAlign='RIGHT')
    elements.append(bottom_wrapper)

    # 생성 일시 (작게)
    elements.append(Spacer(1, 3*mm))
    footer = Paragraph(
        f"생성: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        ParagraphStyle('Footer', parent=styles['Normal'], fontName=font_name,
                      fontSize=7, textColor=colors.grey, alignment=TA_RIGHT)
    )
    elements.append(footer)

    # PDF 생성
    doc.build(elements)


def export_excel_by_range(filepath: str, start_date: str, end_date: str) -> None:
    s = normalize_date(start_date)
    e = normalize_date(end_date)

    conn = connect_db()
    cur = conn.cursor()

    # 전체 헌금 데이터 조회 (헌금종류, 이름, 금액별로 그룹화하여 횟수 계산)
    cur.execute("""
    SELECT o.offering_type, m.name, o.amount, COUNT(*) as count,
           GROUP_CONCAT(COALESCE(o.memo, ''), '|') as memos
    FROM offerings o
    JOIN members m ON m.code = o.member_code
    WHERE o.offering_date BETWEEN ? AND ?
    GROUP BY o.offering_type, m.name, o.amount
    ORDER BY o.offering_type, m.name;
    """, (s, e))
    grouped_data = cur.fetchall()

    # 전체 합계
    cur.execute("""
    SELECT COALESCE(SUM(amount), 0)
    FROM offerings
    WHERE offering_date BETWEEN ? AND ?;
    """, (s, e))
    grand_total = int(cur.fetchone()[0] or 0)

    conn.close()

    # 엑셀 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "헌금내역"

    # 헤더 작성
    ws.append(["예배종류", "헌금종류", "성명", "동반헌금자", "헌금액", "횟수", "비고"])

    # 데이터 작성
    for offering_type, name_val, amount_val, count_val, memos_val in grouped_data:
        # 예배종류 (항상 "주일예배")
        worship_type = "주일예배"

        # 이름 파싱: "이름1(이름2, 이름3)(이름4)" 형식 처리
        import re

        # () 이전까지가 주 헌금자
        main_match = re.match(r'^([^(]+)', str(name_val))
        main_name = main_match.group(1).strip() if main_match else str(name_val)

        # 모든 () 안의 이름들 추출
        companion_matches = re.findall(r'\(([^)]+)\)', str(name_val))

        # 동반헌금자: "주이름, 동반1, 동반2, 동반3"
        if companion_matches:
            # 모든 괄호 안의 내용을 합침
            all_companions = []
            for match in companion_matches:
                # 쉼표로 구분된 이름들을 분리
                names_in_bracket = [n.strip() for n in match.split(',')]
                all_companions.extend(names_in_bracket)

            # 주이름 + 동반자들
            companion_display = main_name + ', ' + ', '.join(all_companions)
        else:
            companion_display = ""

        # 비고 처리 (중복 제거)
        memo_list = [m.strip() for m in str(memos_val).split('|') if m.strip()]
        unique_memos = list(dict.fromkeys(memo_list))  # 순서 유지하며 중복 제거
        memo_display = ', '.join(unique_memos)

        ws.append([
            worship_type,
            offering_type,
            main_name,
            companion_display,
            f"{int(amount_val or 0):,}",
            int(count_val),
            memo_display
        ])

    # 컬럼 너비 조정
    ws.column_dimensions['A'].width = 12  # 예배종류
    ws.column_dimensions['B'].width = 15  # 헌금종류
    ws.column_dimensions['C'].width = 15  # 성명
    ws.column_dimensions['D'].width = 30  # 동반헌금자
    ws.column_dimensions['E'].width = 12  # 헌금액
    ws.column_dimensions['F'].width = 8   # 횟수
    ws.column_dimensions['G'].width = 20  # 비고

    wb.save(filepath)

# =========================
# Date Picker Dialog
# =========================
class DatePicker(tk.Toplevel):
    def __init__(self, master, init_date: str):
        super().__init__(master)
        self.title("날짜 선택")
        self.resizable(False, False)
        self.result = None

        self.transient(master)
        self.grab_set()

        try:
            y, m, d = [int(x) for x in init_date.split("-")]
            self.current = date(y, m, d)
        except Exception:
            self.current = date.today()

        self._build()

        self.update_idletasks()
        x = master.winfo_rootx() + 80
        y = master.winfo_rooty() + 80
        self.geometry(f"+{x}+{y}")

    def _build(self):
        frm = ttk.Frame(self)
        frm.pack(fill="both", expand=True, padx=10, pady=10)

        nav = ttk.Frame(frm)
        nav.pack(fill="x")

        ttk.Button(nav, text="◀", width=3, command=self.prev_month).pack(side="left")
        self.lbl = ttk.Label(nav, text="", width=18, anchor="center")
        self.lbl.pack(side="left", padx=6)
        ttk.Button(nav, text="▶", width=3, command=self.next_month).pack(side="left")
        ttk.Button(nav, text="오늘", width=6, command=self.go_today).pack(side="right", padx=4)

        self.grid_frm = ttk.Frame(frm)
        self.grid_frm.pack(fill="both", expand=True, pady=(8, 0))

        self.render()

    def go_today(self):
        """오늘 날짜로 이동"""
        self.current = date.today()
        self.render()

    def prev_month(self):
        y = self.current.year
        m = self.current.month - 1
        if m <= 0:
            y -= 1
            m = 12
        self.current = date(y, m, 1)
        self.render()

    def next_month(self):
        y = self.current.year
        m = self.current.month + 1
        if m >= 13:
            y += 1
            m = 1
        self.current = date(y, m, 1)
        self.render()

    def render(self):
        for w in self.grid_frm.winfo_children():
            w.destroy()

        y, m = self.current.year, self.current.month
        self.lbl.config(text=f"{y}-{m:02d}")

        days = ["일", "월", "화", "수", "목", "금", "토"]
        head = ttk.Frame(self.grid_frm)
        head.pack(fill="x")
        for i, d in enumerate(days):
            ttk.Label(head, text=d, width=4, anchor="center").grid(row=0, column=i, padx=1, pady=1)

        first = date(y, m, 1)
        start_weekday = (first.weekday() + 1) % 7
        if m == 12:
            next_first = date(y + 1, 1, 1)
        else:
            next_first = date(y, m + 1, 1)
        total_days = (next_first - first).days

        body = ttk.Frame(self.grid_frm)
        body.pack()

        today = date.today()
        r = 0
        c = start_weekday
        for day_num in range(1, total_days + 1):
            def make_cmd(dn=day_num):
                return lambda: self.pick(y, m, dn)

            current_date = date(y, m, day_num)
            btn_text = str(day_num)

            if current_date == today:
                btn_text = f"[{day_num}]"

            b = ttk.Button(body, text=btn_text, width=5, command=make_cmd())
            b.grid(row=r, column=c, padx=1, pady=1)
            c += 1
            if c >= 7:
                c = 0
                r += 1

    def pick(self, y: int, m: int, d: int):
        self.result = f"{y:04d}-{m:02d}-{d:02d}"
        self.destroy()

# =========================
# Member Pick Dialog
# =========================
class MemberPickDialog(tk.Toplevel):
    def __init__(self, master, initial_keyword: str):
        super().__init__(master)
        self.title("사용자 선택")
        self.resizable(False, False)
        self.result = None

        self.transient(master)
        self.grab_set()

        self.var_kw = tk.StringVar(value=initial_keyword.strip())

        wrap = ttk.Frame(self)
        wrap.pack(fill="both", expand=True, padx=10, pady=10)

        top = ttk.Frame(wrap)
        top.pack(fill="x")

        ttk.Label(top, text="이름 검색:").pack(side="left")
        ent = ttk.Entry(top, textvariable=self.var_kw, width=22)
        ent.pack(side="left", padx=6)
        ent.bind("<Return>", lambda _e: self.refresh_list())
        ttk.Button(top, text="재검색", command=self.refresh_list).pack(side="left")

        ttk.Label(wrap, text="(코드/최근헌금일/총액 보고 선택)").pack(anchor="w", pady=(6, 6))

        mid = ttk.Frame(wrap)
        mid.pack(fill="both", expand=True)

        cols = ("code", "name", "last", "total")
        self.tree = ttk.Treeview(mid, columns=cols, show="headings", height=10)
        self.tree.heading("code", text="코드")
        self.tree.heading("name", text="이름")
        self.tree.heading("last", text="최근 헌금일")
        self.tree.heading("total", text="총액(원)")

        self.tree.column("code", width=160, anchor="center")
        self.tree.column("name", width=140, anchor="center")
        self.tree.column("last", width=110, anchor="center")
        self.tree.column("total", width=130, anchor="e")

        vsb = ttk.Scrollbar(mid, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)

        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.tree.bind("<Double-1>", lambda _e: self._select())
        self.tree.bind("<Return>", lambda _e: self._select())

        bot = ttk.Frame(wrap)
        bot.pack(fill="x", pady=(8, 0))

        ttk.Button(bot, text="선택", command=self._select).pack(side="right", padx=4)
        ttk.Button(bot, text="취소(임시코드)", command=self._cancel).pack(side="right", padx=4)

        self.refresh_list()

        self.update_idletasks()
        x = master.winfo_rootx() + 100
        y = master.winfo_rooty() + 100
        self.geometry(f"+{x}+{y}")

        ent.focus_set()
        ent.select_range(0, "end")

    def refresh_list(self):
        kw = self.var_kw.get().strip()
        rows = search_members_with_stats(kw, limit=200)

        for i in self.tree.get_children():
            self.tree.delete(i)

        for code, name, last_date, total_amt in rows:
            self.tree.insert("", "end", values=(code, name, last_date, f"{total_amt:,}"))

        items = self.tree.get_children()
        if items:
            self.tree.selection_set(items[0])
            self.tree.focus(items[0])

    def _select(self):
        sel = self.tree.selection()
        if not sel:
            return
        vals = self.tree.item(sel[0])["values"]
        code = str(vals[0])
        name = str(vals[1])
        self.result = (code, name)
        self.destroy()

    def _cancel(self):
        self.result = None
        self.destroy()

# =========================
# Main Application
# =========================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("교회 헌금 관리 (개선판)")

        # 1366x768 해상도에 맞춤 (작업표시줄 40px 고려)
        window_width = 1340
        window_height = 700
        self.geometry(f"{window_width}x{window_height}")
        self.minsize(1200, 650)

        # 화면 중앙에 배치
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        x = (screen_width - window_width) // 2
        y = (screen_height - window_height) // 2 - 20  # 약간 위쪽으로
        self.geometry(f"{window_width}x{window_height}+{x}+{y}")

        init_db()
        self._build_ui()
        self.refresh_all()

        # 프로그램 시작 시 사용자코드로 포커스
        self.after(100, lambda: self.ent_code.focus_set())

    def pick_date(self, target_var: tk.StringVar, on_done=None):
        init = target_var.get().strip() or datetime.now().strftime("%Y-%m-%d")
        dlg = DatePicker(self, init)
        self.wait_window(dlg)
        if dlg.result:
            target_var.set(dlg.result)
            if callable(on_done):
                on_done()

    @staticmethod
    def _validate_code_digits_only(new_value: str) -> bool:
        if new_value == "":
            return True
        return new_value.isdigit()

    @staticmethod
    def _validate_amount(new_value: str) -> bool:
        """금액 검증: 숫자와 쉼표만, 10억 이하"""
        if new_value == "":
            return True
        cleaned = new_value.replace(",", "")
        if not cleaned.isdigit():
            return False
        return int(cleaned) <= 1_000_000_000

    def _toggle_memo(self):
        """메모 체크박스 토글"""
        if self.var_memo_enabled.get():
            self.ent_memo.config(state="normal")
            self.ent_memo.focus_set()
        else:
            self.ent_memo.config(state="disabled")
            self.var_memo.set("")

    def _format_amount(self, event=None):
        """금액 입력 중 실시간 포맷팅 (타이핑 중에는 포맷 안함)"""
        pass

    def _remove_commas(self, event=None):
        """포커스 들어올 때 쉼표 제거"""
        value = self.var_amount.get().replace(",", "")
        self.var_amount.set(value)
        self.ent_amount.icursor(tk.END)

    def _add_commas(self, event=None):
        """포커스 나갈 때 쉼표 추가"""
        value = self.var_amount.get().replace(",", "")
        if value.isdigit() and value:
            formatted = f"{int(value):,}"
            self.var_amount.set(formatted)

    def _on_name_enter(self):
        """이름 입력 후 Enter 키 처리"""
        self._process_name_input()

    def _on_name_focus_out(self):
        """이름 입력란에서 포커스 아웃 시 처리"""
        # 이미 코드가 있으면 처리 안함
        if self.var_code.get().strip():
            return
        self._process_name_input()

    def _process_name_input(self):
        """이름 입력 처리 로직"""
        name_kw = self.var_name.get().strip()
        if not name_kw:
            return

        # 이미 코드가 있으면 스킵
        if self.var_code.get().strip():
            return

        # 정확히 일치하는 이름 검색
        exact_code = get_member_code_by_name_exact(name_kw)
        if exact_code:
            self.var_code.set(exact_code)
            return

        # 부분 일치 검색
        candidates = search_members_with_stats(name_kw, limit=200)

        if len(candidates) == 0:
            # 일치하는 사용자 없음 - 임시코드 자동 발행
            temp_code = ensure_member_code_by_name(name_kw)
            self.var_code.set(temp_code)
            messagebox.showinfo("안내", f"'{name_kw}' 사용자가 등록되지 않아\n임시 코드 {temp_code}를 발행했습니다.")
        elif len(candidates) == 1:
            # 한 명만 있으면 자동 선택
            code, real_name, _last, _tot = candidates[0]
            self.var_code.set(code)
            self.var_name.set(real_name)
        else:
            # 여러 명 있으면 선택 다이얼로그
            dlg = MemberPickDialog(self, name_kw)
            self.wait_window(dlg)
            if dlg.result:
                code, real_name = dlg.result
                self.var_code.set(code)
                self.var_name.set(real_name)
            else:
                # 취소 시 임시코드 발행
                temp_code = ensure_member_code_by_name(name_kw)
                self.var_code.set(temp_code)
                messagebox.showinfo("안내", f"임시 코드 {temp_code}를 발행했습니다.")

    def _build_ui(self):
        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=8, pady=8)

        self.tab_offerings = ttk.Frame(self.nb)
        self.tab_members = ttk.Frame(self.nb)
        self.tab_types = ttk.Frame(self.nb)
        self.tab_summary = ttk.Frame(self.nb)

        self.nb.add(self.tab_offerings, text="헌금 관리")
        self.nb.add(self.tab_members, text="사용자 관리")
        self.nb.add(self.tab_types, text="헌금 종류 관리")
        self.nb.add(self.tab_summary, text="합계/엑셀/대조")

        self._build_offerings_tab()
        self._build_members_tab()
        self._build_types_tab()
        self._build_summary_tab()

    # =========================
    # Tab 1: Offerings
    # =========================
    def _build_offerings_tab(self):
        frm_top = ttk.LabelFrame(self.tab_offerings, text="입력/수정")
        frm_top.pack(fill="x", padx=8, pady=8)

        self.var_date = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.var_type = tk.StringVar()
        self.var_code = tk.StringVar()
        self.var_name = tk.StringVar()
        self.var_amount = tk.StringVar(value="10,000")  # 기본값 10,000
        self.var_memo = tk.StringVar()
        self.var_memo_enabled = tk.BooleanVar(value=False)  # 메모 체크박스

        row1 = ttk.Frame(frm_top)
        row1.pack(fill="x", padx=8, pady=6)

        ttk.Label(row1, text="날짜").pack(side="left")
        ttk.Entry(row1, textvariable=self.var_date, width=12).pack(side="left", padx=6)
        ttk.Button(row1, text="달력", command=lambda: self.pick_date(self.var_date)).pack(side="left")
        ttk.Button(row1, text="오늘", command=lambda: self.var_date.set(date.today().strftime("%Y-%m-%d"))).pack(side="left", padx=2)

        ttk.Label(row1, text="헌금종류").pack(side="left", padx=(12, 0))
        self.cmb_type = ttk.Combobox(row1, textvariable=self.var_type, values=[], width=16, state="readonly")
        self.cmb_type.pack(side="left", padx=6)

        ttk.Label(row1, text="사용자코드").pack(side="left", padx=(12, 0))
        vcmd = (self.register(self._validate_code_digits_only), "%P")
        self.ent_code = ttk.Entry(row1, textvariable=self.var_code, width=14, validate="key", validatecommand=vcmd)
        self.ent_code.pack(side="left", padx=6)
        self.ent_code.bind("<FocusOut>", lambda _e: self._sync_name_from_code())

        row2 = ttk.Frame(frm_top)
        row2.pack(fill="x", padx=8, pady=6)

        ttk.Label(row2, text="이름").pack(side="left")
        self.ent_name = ttk.Entry(row2, textvariable=self.var_name, width=16)
        self.ent_name.pack(side="left", padx=6)
        self.ent_name.bind("<Return>", lambda _e: self._on_name_enter())
        self.ent_name.bind("<FocusOut>", lambda _e: self._on_name_focus_out())

        ttk.Label(row2, text="금액").pack(side="left", padx=(12, 0))
        vcmd_amt = (self.register(self._validate_amount), "%P")
        self.ent_amount = ttk.Entry(row2, textvariable=self.var_amount, width=14, validate="key", validatecommand=vcmd_amt)
        self.ent_amount.pack(side="left", padx=6)
        self.ent_amount.bind("<KeyRelease>", self._format_amount)
        self.ent_amount.bind("<FocusIn>", self._remove_commas)
        self.ent_amount.bind("<FocusOut>", self._add_commas)
        self.ent_amount.bind("<Return>", self._auto_add_from_amount)
        self.ent_amount.bind("<Tab>", self._auto_add_from_amount)  # Tab 키도 저장
        ttk.Label(row2, text="원").pack(side="left")

        # 메모 체크박스와 입력란
        self.chk_memo = ttk.Checkbutton(row2, text="메모", variable=self.var_memo_enabled,
                                        command=self._toggle_memo, padding=(12, 0, 0, 0))
        self.chk_memo.pack(side="left")
        self.ent_memo = ttk.Entry(row2, textvariable=self.var_memo, width=20, state="disabled")
        self.ent_memo.pack(side="left", padx=6)

        row3 = ttk.Frame(frm_top)
        row3.pack(fill="x", padx=8, pady=(0, 8))

        self.lbl_selected = ttk.Label(row3, text="선택된 항목: (없음)")
        self.lbl_selected.pack(side="left")

        ttk.Button(row3, text="추가 저장", command=self.on_add).pack(side="right", padx=4)
        ttk.Button(row3, text="선택 수정", command=self.on_update).pack(side="right", padx=4)
        ttk.Button(row3, text="선택 삭제", command=self.on_delete).pack(side="right", padx=4)
        ttk.Button(row3, text="다중 삭제", command=self.on_delete_multiple).pack(side="right", padx=4)
        ttk.Button(row3, text="입력 초기화", command=self.clear_form).pack(side="right", padx=4)

        frm_filter = ttk.LabelFrame(self.tab_offerings, text="리스트 조회 기간 및 검색")
        frm_filter.pack(fill="x", padx=8, pady=(0, 8))

        self.list_start = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.list_end = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.filter_name = tk.StringVar()
        self.filter_amount_min = tk.StringVar()
        self.filter_amount_max = tk.StringVar()

        frow1 = ttk.Frame(frm_filter)
        frow1.pack(fill="x", padx=8, pady=6)

        ttk.Label(frow1, text="시작").pack(side="left")
        ttk.Entry(frow1, textvariable=self.list_start, width=12).pack(side="left", padx=6)
        ttk.Button(frow1, text="달력", command=lambda: self.pick_date(self.list_start, self.refresh_offerings)).pack(side="left")

        ttk.Label(frow1, text="종료", padding=(12, 0, 0, 0)).pack(side="left")
        ttk.Entry(frow1, textvariable=self.list_end, width=12).pack(side="left", padx=6)
        ttk.Button(frow1, text="달력", command=lambda: self.pick_date(self.list_end, self.refresh_offerings)).pack(side="left")

        ttk.Label(frow1, text="  빠른설정:", padding=(12, 0, 0, 0)).pack(side="left")
        ttk.Button(frow1, text="오늘", command=self.set_period_today).pack(side="left", padx=2)
        ttk.Button(frow1, text="이번주", command=self.set_period_this_week).pack(side="left", padx=2)
        ttk.Button(frow1, text="이번달", command=self.set_period_this_month).pack(side="left", padx=2)

        frow2 = ttk.Frame(frm_filter)
        frow2.pack(fill="x", padx=8, pady=(0, 6))

        ttk.Label(frow2, text="이름").pack(side="left")
        ent_name = ttk.Entry(frow2, textvariable=self.filter_name, width=12)
        ent_name.pack(side="left", padx=6)
        ent_name.bind("<Return>", lambda e: self.refresh_offerings())

        ttk.Label(frow2, text="금액", padding=(12, 0, 0, 0)).pack(side="left")
        ttk.Entry(frow2, textvariable=self.filter_amount_min, width=10).pack(side="left", padx=6)
        ttk.Label(frow2, text="~").pack(side="left")
        ttk.Entry(frow2, textvariable=self.filter_amount_max, width=10).pack(side="left", padx=6)
        ttk.Label(frow2, text="원").pack(side="left")

        btn_search = ttk.Button(frow2, text="🔍 검색", command=self.refresh_offerings)
        btn_search.pack(side="left", padx=(12, 0))
        btn_search.bind("<Return>", lambda e: self.refresh_offerings())

        btn_clear_filter = ttk.Button(frow2, text="🔄 조건 초기화", command=self.clear_filters)
        btn_clear_filter.pack(side="left", padx=4)
        btn_clear_filter.bind("<Return>", lambda e: self.clear_filters())

        btn_refresh = ttk.Button(frow2, text="🔄 새로고침", command=self.refresh_offerings)
        btn_refresh.pack(side="right")
        btn_refresh.bind("<Return>", lambda e: self.refresh_offerings())

        frm_table = ttk.Frame(self.tab_offerings)
        frm_table.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        cols = ("id", "date", "type", "code", "name", "amount", "memo")
        self.tree = ttk.Treeview(frm_table, columns=cols, show="headings", height=12)
        for c, t in [
            ("id", "ID"), ("date", "날짜"), ("type", "종류"),
            ("code", "코드"), ("name", "이름"), ("amount", "금액(원)"), ("memo", "메모")
        ]:
            self.tree.heading(c, text=t, command=lambda _c=c: self.sort_offerings(_c))

        self.tree.column("id", width=50, anchor="center")
        self.tree.column("date", width=100, anchor="center")
        self.tree.column("type", width=120, anchor="center")
        self.tree.column("code", width=100, anchor="center")
        self.tree.column("name", width=100, anchor="center")
        self.tree.column("amount", width=100, anchor="e")
        self.tree.column("memo", width=180, anchor="w")

        vsb = ttk.Scrollbar(frm_table, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.tree.bind("<<TreeviewSelect>>", self.on_select_row)

        # 정렬 상태 저장
        self.offerings_sort_column = None
        self.offerings_sort_reverse = False

    def set_period_today(self):
        today = date.today().strftime("%Y-%m-%d")
        self.list_start.set(today)
        self.list_end.set(today)
        self.refresh_offerings()

    def set_period_this_week(self):
        today = date.today()
        start = today - timedelta(days=today.weekday() + 1)
        end = start + timedelta(days=6)
        self.list_start.set(start.strftime("%Y-%m-%d"))
        self.list_end.set(end.strftime("%Y-%m-%d"))
        self.refresh_offerings()

    def set_period_this_month(self):
        today = date.today()
        start = date(today.year, today.month, 1)
        if today.month == 12:
            end = date(today.year, 12, 31)
        else:
            end = date(today.year, today.month + 1, 1) - timedelta(days=1)
        self.list_start.set(start.strftime("%Y-%m-%d"))
        self.list_end.set(end.strftime("%Y-%m-%d"))
        self.refresh_offerings()

    def refresh_type_combobox(self):
        types = get_offering_types()
        self.cmb_type["values"] = types
        if types:
            if self.var_type.get().strip() not in types:
                self.var_type.set(types[0])

    def _auto_add_from_amount(self, event):
        """금액 입력 후 Enter/Tab 시 자동 저장"""
        self.on_add()
        self.ent_code.focus_set()  # 사용자코드로 포커스 이동
        return "break"

    def _sync_name_from_code(self):
        code = self.var_code.get().strip()
        if not code:
            return
        name = get_member_name(code)
        if name:
            self.var_name.set(name)

    def _sync_code_from_name_advanced(self):
        """기존 함수 - 더 이상 사용 안함"""
        pass

    def _get_selected_offering_id(self) -> Optional[int]:
        sel = self.tree.selection()
        if not sel:
            return None
        item = self.tree.item(sel[0])
        return int(item["values"][0])

    def on_select_row(self, _evt=None):
        oid = self._get_selected_offering_id()
        if oid is None:
            self.lbl_selected.config(text="선택된 항목: (없음)")
            return
        item = self.tree.item(self.tree.selection()[0])["values"]
        self.lbl_selected.config(text=f"선택된 항목: ID={item[0]}")
        self.var_date.set(item[1])
        self.var_type.set(item[2])
        self.var_code.set(item[3])
        self.var_name.set(item[4])
        self.var_amount.set(str(item[5]).replace(",", ""))
        memo = item[6] if len(item) > 6 else ""
        if memo:
            self.var_memo.set(memo)
            self.var_memo_enabled.set(True)
            self.ent_memo.config(state="normal")
        else:
            self.var_memo.set("")
            self.var_memo_enabled.set(False)
            self.ent_memo.config(state="disabled")

    def _validate_form(self) -> Optional[Tuple[str, str, str, str, int, str]]:
        try:
            date_ = normalize_date(self.var_date.get())
        except Exception:
            messagebox.showerror("오류", "날짜 형식이 올바르지 않아. 예: 2026-01-12")
            return None

        type_ = self.var_type.get().strip()
        if not type_:
            messagebox.showerror("오류", "헌금 종류를 선택해줘.")
            return None

        code = self.var_code.get().strip()
        name = self.var_name.get().strip()

        if not name and not code:
            messagebox.showerror("오류", "이름(또는 사용자코드)을 입력해줘.")
            return None

        # 코드가 있는 경우 DB에 존재하는지 확인
        if code:
            existing_name = get_member_name(code)
            if not existing_name:
                messagebox.showerror("오류", f"사용자 코드 '{code}'가 DB에 등록되어 있지 않습니다.\n사용자 관리 탭에서 먼저 등록해주세요.")
                return None
            # DB에 있는 이름으로 업데이트
            name = existing_name
            self.var_name.set(name)

        # 코드는 없고 이름만 있는 경우 - 임시코드 발행
        if name and not code:
            temp_code = ensure_member_code_by_name(name)
            self.var_code.set(temp_code)
            code = temp_code
            messagebox.showinfo("안내", f"'{name}' 사용자에게 임시 코드 {temp_code}를 발행했습니다.")

        amt_raw = self.var_amount.get().strip().replace(",", "")
        if not amt_raw.isdigit():
            messagebox.showerror("오류", "금액은 0 이상의 숫자로 입력해줘.")
            return None
        amount = int(amt_raw)

        if amount > 1_000_000_000:
            messagebox.showerror("오류", "금액이 너무 커. 10억원 이하로 입력해줘.")
            return None

        memo = self.var_memo.get().strip() if self.var_memo_enabled.get() else ""

        upsert_member(code, name)
        return date_, type_, code, name, amount, memo

    def on_add(self):
        v = self._validate_form()
        if v is None:
            return
        date_, type_, code, name, amount, memo = v

        try:
            insert_offering(date_, type_, code, amount, memo)
        except sqlite3.IntegrityError as e:
            messagebox.showerror("DB 오류", f"저장 실패: {e}")
            return

        self.refresh_type_combobox()
        self.refresh_all()

        # 입력 후 초기화
        self.var_code.set("")
        self.var_name.set("")
        self.var_amount.set("10,000")
        self.var_memo.set("")
        self.var_memo_enabled.set(False)
        self.ent_memo.config(state="disabled")
        self.ent_code.focus_set()  # 사용자코드로 포커스

    def on_update(self):
        oid = self._get_selected_offering_id()
        if oid is None:
            messagebox.showwarning("안내", "수정할 항목을 먼저 선택해줘.")
            return

        v = self._validate_form()
        if v is None:
            return
        date_, type_, code, name, amount, memo = v

        try:
            update_offering(oid, date_, type_, code, amount, memo)
        except sqlite3.IntegrityError as e:
            messagebox.showerror("DB 오류", f"수정 실패: {e}")
            return

        self.refresh_type_combobox()
        self.refresh_all()

    def on_delete(self):
        oid = self._get_selected_offering_id()
        if oid is None:
            messagebox.showwarning("안내", "삭제할 항목을 먼저 선택해줘.")
            return
        if not messagebox.askyesno("확인", f"ID={oid} 항목을 정말 삭제할까?"):
            return
        delete_offering(oid)
        self.refresh_all()
        self.clear_form()

    def on_delete_multiple(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("안내", "삭제할 항목을 선택해줘. (Ctrl+클릭으로 다중선택)")
            return

        count = len(selected)
        if not messagebox.askyesno("확인", f"선택한 {count}개 항목을 정말 삭제할까?"):
            return

        ids = [int(self.tree.item(item)["values"][0]) for item in selected]
        deleted = delete_offerings_batch(ids)

        messagebox.showinfo("완료", f"{deleted}개 항목을 삭제했어.")
        self.refresh_all()
        self.clear_form()

    def clear_form(self):
        self.var_date.set(datetime.now().strftime("%Y-%m-%d"))
        self.var_code.set("")
        self.var_name.set("")
        self.var_amount.set("10,000")  # 기본값으로 리셋
        self.var_memo.set("")
        self.var_memo_enabled.set(False)
        self.ent_memo.config(state="disabled")
        self.lbl_selected.config(text="선택된 항목: (없음)")
        self.tree.selection_remove(self.tree.selection())
        self.ent_code.focus_set()  # 사용자코드로 포커스

    def clear_filters(self):
        """검색 조건 초기화"""
        self.filter_name.set("")
        self.filter_amount_min.set("")
        self.filter_amount_max.set("")
        self.refresh_offerings()

    def refresh_offerings(self):
        for i in self.tree.get_children():
            self.tree.delete(i)

        try:
            start = normalize_date(self.list_start.get())
            end = normalize_date(self.list_end.get())

            if start > end:
                messagebox.showerror("오류", "시작일이 종료일보다 늦어. 날짜를 확인해줘.")
                return

            rows = query_offerings_by_range(start, end)
        except Exception as e:
            messagebox.showerror("오류", f"리스트 기간이 잘못됐어: {e}")
            return

        # 필터 조건 가져오기
        name_filter = self.filter_name.get().strip().lower()
        min_amt = self.filter_amount_min.get().strip().replace(",", "")
        max_amt = self.filter_amount_max.get().strip().replace(",", "")

        min_amount = int(min_amt) if min_amt.isdigit() else None
        max_amount = int(max_amt) if max_amt.isdigit() else None

        # 필터링 및 표시
        for rid, d, t, c, n, a, m in rows:
            # 이름 필터
            if name_filter and name_filter not in n.lower():
                continue

            # 금액 필터
            amount = int(a)
            if min_amount is not None and amount < min_amount:
                continue
            if max_amount is not None and amount > max_amount:
                continue

            self.tree.insert("", "end", values=(rid, d, t, c, n, f"{amount:,}", m))

    def sort_offerings(self, col):
        """헌금 리스트 정렬"""
        # 같은 컬럼 클릭시 오름차순/내림차순 토글
        if self.offerings_sort_column == col:
            self.offerings_sort_reverse = not self.offerings_sort_reverse
        else:
            self.offerings_sort_column = col
            self.offerings_sort_reverse = False

        # 현재 트리의 데이터 가져오기
        data = [(self.tree.set(item, col), item) for item in self.tree.get_children("")]

        # 정렬 (금액은 숫자로, 나머지는 문자열로)
        if col == "amount":
            data.sort(key=lambda x: int(x[0].replace(",", "")), reverse=self.offerings_sort_reverse)
        elif col == "id":
            data.sort(key=lambda x: int(x[0]), reverse=self.offerings_sort_reverse)
        else:
            data.sort(reverse=self.offerings_sort_reverse)

        # 정렬된 순서대로 재배치
        for index, (val, item) in enumerate(data):
            self.tree.move(item, "", index)

        # 헤더에 정렬 표시
        for c in ("id", "date", "type", "code", "name", "amount", "memo"):
            if c == col:
                direction = "▼" if self.offerings_sort_reverse else "▲"
                self.tree.heading(c, text=f"{self._get_col_name(c)} {direction}")
            else:
                self.tree.heading(c, text=self._get_col_name(c))

    def _get_col_name(self, col):
        """컬럼 한글 이름 반환"""
        names = {
            "id": "ID", "date": "날짜", "type": "종류",
            "code": "코드", "name": "이름", "amount": "금액(원)", "memo": "메모"
        }
        return names.get(col, col)

    # =========================
    # Tab 2: Members
    # =========================
    def _build_members_tab(self):
        outer = ttk.Frame(self.tab_members)
        outer.pack(fill="both", expand=True, padx=8, pady=8)

        left = ttk.LabelFrame(outer, text="사용자 등록/수정")
        left.pack(side="left", fill="y", padx=(0, 8))

        self.m_code = tk.StringVar()
        self.m_name = tk.StringVar()

        row = ttk.Frame(left)
        row.pack(padx=10, pady=10)

        ttk.Label(row, text="코드(숫자만)").grid(row=0, column=0, sticky="w")
        vcmd2 = (self.register(self._validate_code_digits_only), "%P")
        ttk.Entry(row, textvariable=self.m_code, width=18, validate="key", validatecommand=vcmd2).grid(row=0, column=1, padx=6, pady=4)

        ttk.Label(row, text="이름").grid(row=1, column=0, sticky="w")
        ttk.Entry(row, textvariable=self.m_name, width=18).grid(row=1, column=1, padx=6, pady=4)

        btns = ttk.Frame(left)
        btns.pack(fill="x", padx=10, pady=(0, 10))

        btn_save = ttk.Button(btns, text="💾 저장(등록/수정)", command=self.on_member_save, style='Primary.TButton')
        btn_save.pack(fill="x", pady=4)
        btn_save.bind("<Return>", lambda e: self.on_member_save())

        btn_del = ttk.Button(btns, text="🗑️ 선택 삭제", command=self.on_member_delete, style='Danger.TButton')
        btn_del.pack(fill="x", pady=4)
        btn_del.bind("<Return>", lambda e: self.on_member_delete())

        btn_clear = ttk.Button(btns, text="🔄 입력 초기화", command=self.on_member_clear)
        btn_clear.pack(fill="x", pady=4)
        btn_clear.bind("<Return>", lambda e: self.on_member_clear())

        btn_import = ttk.Button(btns, text="📁 엑셀로 일괄 등록", command=self.on_member_import_excel)
        btn_import.pack(fill="x", pady=4)
        btn_import.bind("<Return>", lambda e: self.on_member_import_excel())

        right = ttk.LabelFrame(outer, text="사용자 목록")
        right.pack(side="left", fill="both", expand=True)

        top = ttk.Frame(right)
        top.pack(fill="x", padx=8, pady=8)

        self.m_search = tk.StringVar()
        ttk.Label(top, text="검색:").pack(side="left")
        ent = ttk.Entry(top, textvariable=self.m_search, width=22)
        ent.pack(side="left", padx=6)
        ent.bind("<Return>", lambda _e: self.refresh_members())

        btn_search = ttk.Button(top, text="🔍 검색", command=self.refresh_members)
        btn_search.pack(side="left")
        btn_search.bind("<Return>", lambda e: self.refresh_members())

        btn_all = ttk.Button(top, text="📋 전체", command=lambda: (self.m_search.set(""), self.refresh_members()))
        btn_all.pack(side="left", padx=6)
        btn_all.bind("<Return>", lambda e: (self.m_search.set(""), self.refresh_members()))

        mid = ttk.Frame(right)
        mid.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        self.member_tree = ttk.Treeview(mid, columns=("code", "name"), show="headings", height=18)
        self.member_tree.heading("code", text="코드", command=lambda: self.sort_members("code"))
        self.member_tree.heading("name", text="이름", command=lambda: self.sort_members("name"))
        self.member_tree.column("code", width=200, anchor="center")
        self.member_tree.column("name", width=200, anchor="center")

        vsb = ttk.Scrollbar(mid, orient="vertical", command=self.member_tree.yview)
        self.member_tree.configure(yscrollcommand=vsb.set)

        self.member_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.member_tree.bind("<<TreeviewSelect>>", self.on_member_select)

        # 정렬 상태
        self.members_sort_column = None
        self.members_sort_reverse = False

    def on_member_clear(self):
        self.m_code.set("")
        self.m_name.set("")
        self.member_tree.selection_remove(self.member_tree.selection())

    def on_member_select(self, _evt=None):
        sel = self.member_tree.selection()
        if not sel:
            return
        vals = self.member_tree.item(sel[0])["values"]
        self.m_code.set(vals[0])
        self.m_name.set(vals[1])

    def on_member_save(self):
        code = self.m_code.get().strip()
        name = self.m_name.get().strip()
        if not code or not name:
            messagebox.showerror("오류", "코드/이름을 둘 다 입력해줘.")
            return
        if not code.isdigit():
            messagebox.showerror("오류", "사용자 코드는 숫자만 가능해.")
            return
        upsert_member(code, name)
        self.refresh_all()

    def on_member_delete(self):
        sel = self.member_tree.selection()
        if not sel:
            messagebox.showwarning("안내", "삭제할 사용자를 선택해줘.")
            return
        code = self.member_tree.item(sel[0])["values"][0]
        if not messagebox.askyesno("확인", f"사용자 코드 {code}를 삭제할까?\n(헌금 데이터가 있으면 삭제 불가)"):
            return
        try:
            delete_member(code)
        except sqlite3.IntegrityError:
            messagebox.showerror("오류", "이 사용자는 헌금 내역이 있어서 삭제할 수 없어.")
            return
        self.refresh_all()
        self.on_member_clear()

    def on_member_import_excel(self):
        filepath = filedialog.askopenfilename(
            title="사용자 엑셀 선택",
            filetypes=[("Excel 파일", "*.xlsx")]
        )
        if not filepath:
            return
        try:
            ok, skip = import_members_from_excel(filepath)
        except Exception as e:
            messagebox.showerror("오류", f"엑셀 불러오기 실패: {e}")
            return
        messagebox.showinfo("완료", f"일괄 등록 완료!\n성공: {ok}건\n스킵: {skip}건")
        self.refresh_all()

    def refresh_members(self):
        for i in self.member_tree.get_children():
            self.member_tree.delete(i)
        rows = query_members(self.m_search.get())
        for code, name in rows:
            self.member_tree.insert("", "end", values=(code, name))

    def sort_members(self, col):
        """사용자 리스트 정렬"""
        if self.members_sort_column == col:
            self.members_sort_reverse = not self.members_sort_reverse
        else:
            self.members_sort_column = col
            self.members_sort_reverse = False

        data = [(self.member_tree.set(item, col), item) for item in self.member_tree.get_children("")]
        data.sort(reverse=self.members_sort_reverse)

        for index, (val, item) in enumerate(data):
            self.member_tree.move(item, "", index)

        # 헤더 표시
        for c in ("code", "name"):
            if c == col:
                direction = "▼" if self.members_sort_reverse else "▲"
                col_name = "코드" if c == "code" else "이름"
                self.member_tree.heading(c, text=f"{col_name} {direction}")
            else:
                col_name = "코드" if c == "code" else "이름"
                self.member_tree.heading(c, text=col_name)

    # =========================
    # Tab 3: Types
    # =========================
    def _build_types_tab(self):
        outer = ttk.Frame(self.tab_types)
        outer.pack(fill="both", expand=True, padx=8, pady=8)

        left = ttk.LabelFrame(outer, text="헌금 종류 관리")
        left.pack(side="left", fill="y", padx=(0, 8))

        self.t_name = tk.StringVar()
        self.t_new_name = tk.StringVar()

        box = ttk.Frame(left)
        box.pack(padx=10, pady=10)

        ttk.Label(box, text="선택된 종류").grid(row=0, column=0, sticky="w")
        ttk.Entry(box, textvariable=self.t_name, width=20).grid(row=0, column=1, padx=6, pady=4)

        ttk.Label(box, text="새 이름(변경)").grid(row=1, column=0, sticky="w")
        ttk.Entry(box, textvariable=self.t_new_name, width=20).grid(row=1, column=1, padx=6, pady=4)

        btns = ttk.Frame(left)
        btns.pack(fill="x", padx=10, pady=(0, 10))

        btn_add = ttk.Button(btns, text="➕ 추가", command=self.on_type_add, style='Primary.TButton')
        btn_add.pack(fill="x", pady=4)
        btn_add.bind("<Return>", lambda e: self.on_type_add())

        btn_rename = ttk.Button(btns, text="✏️ 이름 변경", command=self.on_type_rename)
        btn_rename.pack(fill="x", pady=4)
        btn_rename.bind("<Return>", lambda e: self.on_type_rename())

        btn_delete = ttk.Button(btns, text="🗑️ 삭제", command=self.on_type_delete, style='Danger.TButton')
        btn_delete.pack(fill="x", pady=4)
        btn_delete.bind("<Return>", lambda e: self.on_type_delete())

        btn_clear = ttk.Button(btns, text="🔄 초기화", command=self.on_type_clear)
        btn_clear.pack(fill="x", pady=4)
        btn_clear.bind("<Return>", lambda e: self.on_type_clear())

        right = ttk.LabelFrame(outer, text="헌금 종류 목록")
        right.pack(side="left", fill="both", expand=True)

        mid = ttk.Frame(right)
        mid.pack(fill="both", expand=True, padx=8, pady=8)

        self.type_tree = ttk.Treeview(mid, columns=("name",), show="headings", height=18)
        self.type_tree.heading("name", text="헌금 종류")
        self.type_tree.column("name", width=240, anchor="center")

        vsb = ttk.Scrollbar(mid, orient="vertical", command=self.type_tree.yview)
        self.type_tree.configure(yscrollcommand=vsb.set)

        self.type_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.type_tree.bind("<<TreeviewSelect>>", self.on_type_select)

    def refresh_types(self):
        for i in self.type_tree.get_children():
            self.type_tree.delete(i)
        for t in get_offering_types():
            self.type_tree.insert("", "end", values=(t,))
        self.refresh_type_combobox()

    def on_type_clear(self):
        self.t_name.set("")
        self.t_new_name.set("")
        self.type_tree.selection_remove(self.type_tree.selection())

    def on_type_select(self, _evt=None):
        sel = self.type_tree.selection()
        if not sel:
            return
        name = self.type_tree.item(sel[0])["values"][0]
        self.t_name.set(name)
        self.t_new_name.set("")

    def on_type_add(self):
        name = self.t_name.get().strip()
        if not name:
            messagebox.showerror("오류", "추가할 헌금 종류 이름을 입력해줘.")
            return
        try:
            add_offering_type(name)
        except sqlite3.IntegrityError:
            messagebox.showerror("오류", "이미 존재하는 헌금 종류야.")
            return
        self.refresh_all()
        self.on_type_clear()

    def on_type_rename(self):
        old = self.t_name.get().strip()
        new = self.t_new_name.get().strip()
        if not old or not new:
            messagebox.showerror("오류", "기존 종류와 새 이름을 입력해줘.")
            return
        try:
            rename_offering_type(old, new)
        except Exception as e:
            messagebox.showerror("오류", str(e))
            return
        self.refresh_all()
        self.on_type_clear()

    def on_type_delete(self):
        name = self.t_name.get().strip()
        if not name:
            messagebox.showwarning("안내", "삭제할 헌금 종류를 선택해줘.")
            return
        if not messagebox.askyesno("확인", f"헌금 종류 [{name}]를 삭제할까?\n(헌금 내역이 있으면 삭제 불가)"):
            return
        try:
            delete_offering_type(name)
        except Exception as e:
            messagebox.showerror("오류", str(e))
            return
        self.refresh_all()
        self.on_type_clear()

    # =========================
    # Tab 4: Summary
    # =========================
    def _build_summary_tab(self):
        outer = ttk.Frame(self.tab_summary)
        outer.pack(fill="both", expand=True, padx=8, pady=8)

        # 상단 프레임: 기간 선택 + 버튼들
        top_frame = ttk.Frame(outer)
        top_frame.pack(fill="x", padx=8, pady=8)

        period = ttk.LabelFrame(top_frame, text="기간 선택")
        period.pack(side="left", fill="x", expand=True, padx=(0, 8))

        self.s_start = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.s_end = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))

        pr = ttk.Frame(period)
        pr.pack(fill="x", padx=8, pady=8)

        ttk.Label(pr, text="시작").pack(side="left")
        ttk.Entry(pr, textvariable=self.s_start, width=12).pack(side="left", padx=6)
        ttk.Button(pr, text="달력", command=lambda: self.pick_date(self.s_start, self.refresh_summary)).pack(side="left")

        ttk.Label(pr, text="종료", padding=(12, 0, 0, 0)).pack(side="left")
        ttk.Entry(pr, textvariable=self.s_end, width=12).pack(side="left", padx=6)
        ttk.Button(pr, text="달력", command=lambda: self.pick_date(self.s_end, self.refresh_summary)).pack(side="left")

        ttk.Label(pr, text="  빠른설정:", padding=(12, 0, 0, 0)).pack(side="left")
        ttk.Button(pr, text="오늘", command=self.set_summary_today).pack(side="left", padx=2)
        ttk.Button(pr, text="이번주", command=self.set_summary_week).pack(side="left", padx=2)
        ttk.Button(pr, text="이번달", command=self.set_summary_month).pack(side="left", padx=2)

        ttk.Button(pr, text="새로고침", command=self.refresh_summary).pack(side="right")

        # 오른쪽 버튼 그룹
        btn_group = ttk.LabelFrame(top_frame, text="엑셀/백업")
        btn_group.pack(side="right", fill="y")

        btn_inner = ttk.Frame(btn_group)
        btn_inner.pack(padx=10, pady=10)

        btn_backup = ttk.Button(btn_inner, text="💾 DB 백업", command=self.on_backup_db, width=12)
        btn_backup.pack(pady=4)
        btn_backup.bind("<Return>", lambda e: self.on_backup_db())

        btn_excel = ttk.Button(btn_inner, text="📊 엑셀 저장", command=self.on_export_excel, width=12, style='Primary.TButton')
        btn_excel.pack(pady=4)
        btn_excel.bind("<Return>", lambda e: self.on_export_excel())

        btn_pdf = ttk.Button(btn_inner, text="📄 PDF 리포트", command=self.on_export_pdf, width=12, style='Primary.TButton')
        btn_pdf.pack(pady=4)
        btn_pdf.bind("<Return>", lambda e: self.on_export_pdf())

        totalbox = ttk.LabelFrame(outer, text="기간 전체 합계")
        totalbox.pack(fill="x", padx=8, pady=(0, 8))

        self.lbl_period_total = ttk.Label(totalbox, text="0원", font=("맑은 고딕", 12, "bold"))
        self.lbl_period_total.pack(anchor="w", padx=10, pady=10)

        typebox = ttk.LabelFrame(outer, text="헌금 종류별 합계")
        typebox.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        mid = ttk.Frame(typebox)
        mid.pack(fill="both", expand=True, padx=8, pady=8)

        self.sum_tree = ttk.Treeview(mid, columns=("type", "total"), show="headings", height=6)
        self.sum_tree.heading("type", text="헌금 종류", command=lambda: self.sort_summary("type"))
        self.sum_tree.heading("total", text="합계(원)", command=lambda: self.sort_summary("total"))
        self.sum_tree.column("type", width=200, anchor="center")
        self.sum_tree.column("total", width=180, anchor="e")

        vsb = ttk.Scrollbar(mid, orient="vertical", command=self.sum_tree.yview)
        self.sum_tree.configure(yscrollcommand=vsb.set)

        self.sum_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        # 정렬 상태
        self.summary_sort_column = None
        self.summary_sort_reverse = False

        check = ttk.LabelFrame(outer, text="전체금액 대조")
        check.pack(fill="x", padx=8, pady=(0, 8))
        self._build_check_ui(check)

    def set_summary_today(self):
        today = date.today().strftime("%Y-%m-%d")
        self.s_start.set(today)
        self.s_end.set(today)
        self.refresh_summary()

    def set_summary_week(self):
        today = date.today()
        start = today - timedelta(days=today.weekday() + 1)
        end = start + timedelta(days=6)
        self.s_start.set(start.strftime("%Y-%m-%d"))
        self.s_end.set(end.strftime("%Y-%m-%d"))
        self.refresh_summary()

    def set_summary_month(self):
        today = date.today()
        start = date(today.year, today.month, 1)
        if today.month == 12:
            end = date(today.year, 12, 31)
        else:
            end = date(today.year, today.month + 1, 1) - timedelta(days=1)
        self.s_start.set(start.strftime("%Y-%m-%d"))
        self.s_end.set(end.strftime("%Y-%m-%d"))
        self.refresh_summary()

    def _build_check_ui(self, parent):
        frm = ttk.Frame(parent)
        frm.pack(fill="x", padx=10, pady=10)

        ttk.Label(frm, text="항목").grid(row=0, column=0, padx=4, pady=3, sticky="w")
        ttk.Label(frm, text="단가").grid(row=0, column=1, padx=4, pady=3, sticky="w")
        ttk.Label(frm, text="수량").grid(row=0, column=2, padx=4, pady=3, sticky="w")
        ttk.Label(frm, text="입력금액").grid(row=0, column=3, padx=4, pady=3, sticky="w")
        ttk.Label(frm, text="합계금액").grid(row=0, column=4, padx=4, pady=3, sticky="w")

        self.check_rows = []
        items = [
            ("5만원권", 50_000),
            ("1만원권", 10_000),
            ("5천원권", 5_000),
            ("1천원권", 1_000),
            ("동전", 0),
            ("수표 100만원권", 1_000_000),
            ("수표 10만원권", 100_000),
            ("기타", 0),
        ]

        for i, (name, unit) in enumerate(items, start=1):
            v_qty = tk.StringVar(value="")
            v_amt = tk.StringVar(value="")

            ttk.Label(frm, text=name).grid(row=i, column=0, padx=4, pady=3, sticky="w")
            ttk.Label(frm, text=f"{unit:,}").grid(row=i, column=1, padx=4, pady=3, sticky="e")

            e_qty = ttk.Entry(frm, textvariable=v_qty, width=10)
            e_qty.grid(row=i, column=2, padx=4, pady=3)
            e_amt = ttk.Entry(frm, textvariable=v_amt, width=14)
            e_amt.grid(row=i, column=3, padx=4, pady=3)

            lbl_sum = ttk.Label(frm, text="0원")
            lbl_sum.grid(row=i, column=4, padx=6, pady=3, sticky="w")

            self.check_rows.append((name, unit, v_qty, v_amt, lbl_sum))

            v_qty.trace_add("write", lambda *_: self.calc_check())
            v_amt.trace_add("write", lambda *_: self.calc_check())

        r = len(items) + 1
        ttk.Separator(frm, orient="horizontal").grid(row=r, column=0, columnspan=5, sticky="ew", pady=6)

        r += 1
        ttk.Label(frm, text="대조 합계", font=("맑은 고딕", 11, "bold")).grid(row=r, column=0, padx=4, pady=3, sticky="w")
        self.lbl_check_total = ttk.Label(frm, text="0원", font=("맑은 고딕", 12, "bold"), foreground="#27AE60")
        self.lbl_check_total.grid(row=r, column=4, padx=4, pady=3, sticky="w")

        # 지출내역 추가
        r += 1
        ttk.Label(frm, text="지출내역", font=("맑은 고딕", 11, "bold"), foreground="#E74C3C").grid(row=r, column=0, padx=4, pady=3, sticky="w")
        ttk.Label(frm, text="금액:").grid(row=r, column=1, padx=4, pady=3, sticky="e")
        self.var_expense = tk.StringVar(value="")
        e_expense = ttk.Entry(frm, textvariable=self.var_expense, width=10)
        e_expense.grid(row=r, column=2, padx=4, pady=3)

        ttk.Label(frm, text="사용처:").grid(row=r, column=3, padx=4, pady=3, sticky="w")
        self.lbl_expense = ttk.Label(frm, text="0원", font=("맑은 고딕", 12), foreground="#E74C3C")
        self.lbl_expense.grid(row=r, column=4, padx=4, pady=3, sticky="w")
        self.var_expense.trace_add("write", lambda *_: self.calc_check())

        # 지출 메모 (다음 줄)
        r += 1
        ttk.Label(frm, text="").grid(row=r, column=0, padx=4, pady=3)
        ttk.Label(frm, text="메모:").grid(row=r, column=1, padx=4, pady=3, sticky="e")
        self.var_expense_memo = tk.StringVar(value="")
        e_expense_memo = ttk.Entry(frm, textvariable=self.var_expense_memo, width=45)
        e_expense_memo.grid(row=r, column=2, columnspan=3, padx=4, pady=3, sticky="ew")

        # 실제 잔액 (대조합계 - 지출)
        r += 1
        ttk.Label(frm, text="실제 잔액 (대조-지출)", font=("맑은 고딕", 11, "bold")).grid(row=r, column=0, padx=4, pady=3, sticky="w")
        self.lbl_actual_balance = ttk.Label(frm, text="0원", font=("맑은 고딕", 12, "bold"), foreground="#2980B9")
        self.lbl_actual_balance.grid(row=r, column=4, padx=4, pady=3, sticky="w")

        r += 1
        ttk.Separator(frm, orient="horizontal").grid(row=r, column=0, columnspan=5, sticky="ew", pady=6)

        r += 1
        ttk.Label(frm, text="차이 (헌금합계-실제잔액)", font=("맑은 고딕", 11, "bold")).grid(row=r, column=0, padx=4, pady=3, sticky="w")
        self.lbl_check_diff = ttk.Label(frm, text="0원", font=("맑은 고딕", 12, "bold"), foreground="#C0392B")
        self.lbl_check_diff.grid(row=r, column=4, padx=4, pady=3, sticky="w")

        r += 1
        btn_clear = ttk.Button(frm, text="🔄 초기화", command=self.clear_check)
        btn_clear.grid(row=r, column=3, padx=4, pady=6, sticky="e")
        btn_clear.bind("<Return>", lambda e: self.clear_check())

        btn_calc = ttk.Button(frm, text="🧮 계산", command=self.calc_check, style='Primary.TButton')
        btn_calc.grid(row=r, column=4, padx=4, pady=6, sticky="e")
        btn_calc.bind("<Return>", lambda e: self.calc_check())

    def _safe_int(self, s: str) -> int:
        s = (s or "").strip().replace(",", "")
        return int(s) if s.isdigit() else 0

    def calc_check(self):
        try:
            total = sum_total_by_range(self.s_start.get(), self.s_end.get())
        except Exception:
            total = 0

        check_total = 0

        for name, unit, v_qty, v_amt, lbl_sum in self.check_rows:
            qty = self._safe_int(v_qty.get())
            amt = self._safe_int(v_amt.get())

            if unit > 0:
                part = qty * unit if qty > 0 else amt
            else:
                part = amt

            check_total += part
            lbl_sum.config(text=f"{part:,}원")

        # 지출 금액 처리
        expense = self._safe_int(self.var_expense.get())
        self.lbl_expense.config(text=f"{expense:,}원")

        # 실제 잔액 = 대조 합계 - 지출
        actual_balance = check_total - expense

        # 차이 = 헌금 합계 - 실제 잔액
        diff = total - actual_balance

        self.lbl_check_total.config(text=f"{check_total:,}원")
        self.lbl_actual_balance.config(text=f"{actual_balance:,}원")
        self.lbl_check_diff.config(text=f"{diff:,}원")

    def clear_check(self):
        for name, unit, v_qty, v_amt, lbl_sum in self.check_rows:
            v_qty.set("")
            v_amt.set("")
            lbl_sum.config(text="0원")
        self.var_expense.set("")
        self.var_expense_memo.set("")
        self.lbl_expense.config(text="0원")
        self.calc_check()

    def refresh_summary(self):
        try:
            start = normalize_date(self.s_start.get())
            end = normalize_date(self.s_end.get())

            if start > end:
                messagebox.showerror("오류", "시작일이 종료일보다 늦어.")
                return

            total = sum_total_by_range(start, end)
            by_type = sum_by_type_by_range(start, end)
        except Exception as e:
            messagebox.showerror("오류", f"기간이 올바르지 않아: {e}")
            return

        self.lbl_period_total.config(text=f"{total:,}원")

        for i in self.sum_tree.get_children():
            self.sum_tree.delete(i)
        for t, v in by_type.items():
            self.sum_tree.insert("", "end", values=(t, f"{v:,}"))

        self.calc_check()

    def on_export_excel(self):
        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 파일", "*.xlsx")],
            title="엑셀로 저장"
        )
        if not filepath:
            return
        try:
            start = normalize_date(self.s_start.get())
            end = normalize_date(self.s_end.get())
            if start > end:
                messagebox.showerror("오류", "시작일이 종료일보다 늦어.")
                return
            export_excel_by_range(filepath, start, end)
        except Exception as e:
            messagebox.showerror("오류", f"엑셀 저장 실패: {e}")
            return
        messagebox.showinfo("완료", f"엑셀 저장 완료!\n{filepath}")

    def on_backup_db(self):
        try:
            backup_path = backup_database()
            messagebox.showinfo("완료", f"백업 완료!\n{backup_path}")
        except Exception as e:
            messagebox.showerror("오류", f"백업 실패: {e}")

    def on_export_pdf(self):
        """PDF 리포트 출력"""
        if not PDF_AVAILABLE:
            messagebox.showerror("오류",
                "PDF 생성 기능을 사용하려면 reportlab 라이브러리가 필요합니다.\n\n"
                "설치 방법:\n"
                "pip install reportlab")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF 파일", "*.pdf")],
            title="PDF 리포트 저장"
        )
        if not filepath:
            return

        try:
            start = normalize_date(self.s_start.get())
            end = normalize_date(self.s_end.get())
            if start > end:
                messagebox.showerror("오류", "시작일이 종료일보다 늦어.")
                return

            # 대조 데이터 수집
            check_data = {}
            for name, unit, v_qty, v_amt, lbl_sum in self.check_rows:
                qty = self._safe_int(v_qty.get())
                amt = self._safe_int(v_amt.get())
                check_data[name] = {
                    'unit': unit,
                    'qty': qty,
                    'amt': amt
                }

            # 지출 데이터
            expense_amt = self._safe_int(self.var_expense.get())
            expense_memo = self.var_expense_memo.get().strip()

            export_pdf_report(filepath, start, end, check_data, expense_amt, expense_memo)
            messagebox.showinfo("완료", f"PDF 리포트 생성 완료!\n{filepath}")
        except Exception as e:
            messagebox.showerror("오류", f"PDF 생성 실패: {e}")

    def sort_summary(self, col):
        """합계 리스트 정렬"""
        if self.summary_sort_column == col:
            self.summary_sort_reverse = not self.summary_sort_reverse
        else:
            self.summary_sort_column = col
            self.summary_sort_reverse = False

        data = [(self.sum_tree.set(item, col), item) for item in self.sum_tree.get_children("")]

        # 합계는 숫자로 정렬
        if col == "total":
            data.sort(key=lambda x: int(x[0].replace(",", "")), reverse=self.summary_sort_reverse)
        else:
            data.sort(reverse=self.summary_sort_reverse)

        for index, (val, item) in enumerate(data):
            self.sum_tree.move(item, "", index)

        # 헤더 표시
        for c in ("type", "total"):
            if c == col:
                direction = "▼" if self.summary_sort_reverse else "▲"
                col_name = "헌금 종류" if c == "type" else "합계(원)"
                self.sum_tree.heading(c, text=f"{col_name} {direction}")
            else:
                col_name = "헌금 종류" if c == "type" else "합계(원)"
                self.sum_tree.heading(c, text=col_name)

    # =========================
    # Refresh All
    # =========================
    def refresh_all(self):
        self.refresh_types()
        self.refresh_members()
        self.refresh_type_combobox()
        self.refresh_offerings()
        self.refresh_summary()


if __name__ == "__main__":
    app = App()
    app.mainloop()
